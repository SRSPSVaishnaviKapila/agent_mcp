"""
MCP PowerPoint server — rich themed slides.
Exposes MCP tools used by the agent to create beautifully designed presentations.
"""

from __future__ import annotations

import io
import json
import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
import random

import requests
from PIL import Image as PILImage, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from utils.planner import build_fallback_plan

if sys.version_info >= (3, 14):
    raise RuntimeError(
        "MCP server is not supported on Python 3.14. Use Python 3.11 or 3.12."
    )

from mcp.server.fastmcp import FastMCP

mcp = FastMCP("auto-ppt-server")

# ── Palettes ───────────────────────────────────────────────────────────────────

PALETTES = [
    {"bg_dark": "1E2761", "bg_light": "F0F4FF", "accent": "4F8EF7", "text_dark": "FFFFFF", "text_light": "1E2761", "shape": "CADCFC"},
    {"bg_dark": "1B3A2F", "bg_light": "F4F9F1", "accent": "4CAF82", "text_dark": "FFFFFF", "text_light": "1B3A2F", "shape": "A8D5B5"},
    {"bg_dark": "2F3C7E", "bg_light": "FFF8F0", "accent": "F96167", "text_dark": "FFFFFF", "text_light": "2F3C7E", "shape": "F9C96B"},
    {"bg_dark": "3D1C12", "bg_light": "F9F4EE", "accent": "C45E3E", "text_dark": "FFFFFF", "text_light": "3D1C12", "shape": "E8C4A0"},
    {"bg_dark": "021B2C", "bg_light": "F0FAFC", "accent": "02C39A", "text_dark": "FFFFFF", "text_light": "021B2C", "shape": "7DCFC4"},
    {"bg_dark": "3A1028", "bg_light": "FDF5F0", "accent": "C45B8A", "text_dark": "FFFFFF", "text_light": "3A1028", "shape": "F0B8D4"},
]

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


@dataclass
class PresentationState:
    presentation: Presentation | None = None
    slides_added: int = 0
    slide_index: int = 0
    palette: dict = field(default_factory=lambda: random.choice(PALETTES))


STATE = PresentationState()


# ── Drawing helpers ────────────────────────────────────────────────────────────

def _rgb(hex_str: str) -> RGBColor:
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _add_rect(slide, left, top, width, height, hex_color: str):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(hex_color)
    shape.line.color.rgb = _rgb(hex_color)
    return shape


def _add_text_box(slide, text, left, top, width, height, font_name, font_size,
                  bold, hex_color, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(hex_color)
    return txBox


def _extract_json(text: str) -> Any:
    if not isinstance(text, str):
        return None
    clean = re.sub(r"```[a-zA-Z]*\n?", "", text).replace("```", "").strip()
    try:
        return json.loads(clean)
    except Exception:
        pass
    for opener, closer in (("[", "]"), ("{", "}")):
        s, e = clean.find(opener), clean.rfind(closer)
        if s != -1 and e > s:
            try:
                return json.loads(clean[s: e + 1])
            except Exception:
                pass
    return None


def _hf_headers() -> dict[str, str]:
    token = os.getenv("HF_API_TOKEN", "").strip()
    return {"Authorization": f"Bearer {token}"} if token else {}


def _call_hf_text_model(prompt: str) -> str | None:
    model = os.getenv("HF_TEXT_MODEL", "mistralai/Mistral-7B-Instruct-v0.3")
    url = f"https://api-inference.huggingface.co/models/{model}"
    payload = {"inputs": prompt,
               "parameters": {"max_new_tokens": 350, "temperature": 0.5, "return_full_text": False}}
    try:
        r = requests.post(url, headers=_hf_headers(), json=payload, timeout=60)
        r.raise_for_status()
        data = r.json()
    except Exception:
        return None
    if isinstance(data, list) and data and isinstance(data[0], dict):
        return data[0].get("generated_text") or data[0].get("summary_text")
    if isinstance(data, dict):
        return data.get("generated_text")
    return None


def _call_hf_image_model(prompt: str) -> bytes | None:
    model = os.getenv("HF_IMG_MODEL", "stabilityai/stable-diffusion-xl-base-1.0")
    url = f"https://api-inference.huggingface.co/models/{model}"
    try:
        r = requests.post(url, headers=_hf_headers(), json={"inputs": prompt}, timeout=120)
        r.raise_for_status()
    except Exception:
        return None
    if "image" in r.headers.get("Content-Type", "").lower():
        return r.content
    return None


def _fit_image(image_bytes: bytes) -> io.BytesIO:
    img = PILImage.open(io.BytesIO(image_bytes)).convert("RGB")
    fitted = ImageOps.fit(img, (1600, 900), method=PILImage.Resampling.LANCZOS)
    out = io.BytesIO()
    fitted.save(out, format="PNG")
    out.seek(0)
    return out


def _require_presentation() -> Presentation:
    if STATE.presentation is None:
        raise RuntimeError("No active presentation. Call create_presentation first.")
    return STATE.presentation


def _normalize_bullets(bullets: list | None) -> list[str]:
    if not bullets:
        return []
    return [str(b).strip() for b in bullets if str(b).strip()][:7]


def _output_path(filename: str) -> Path:
    out_dir = Path(os.getenv("PPT_OUTPUT_DIR", "outputs").strip() or "outputs")
    out_dir.mkdir(parents=True, exist_ok=True)
    name = Path(filename).name.strip() or "output.pptx"
    if not name.lower().endswith(".pptx"):
        name = f"{name}.pptx"
    return out_dir / name


# ── Slide layout builders ──────────────────────────────────────────────────────

def _build_title_slide(slide, title: str):
    p = STATE.palette
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, p["bg_dark"])
    _add_rect(slide, 0, 0, Inches(0.18), SLIDE_H, p["accent"])
    circ = slide.shapes.add_shape(9, Inches(9.5), Inches(-1.5), Inches(5), Inches(5))
    circ.fill.solid(); circ.fill.fore_color.rgb = _rgb(p["accent"]); circ.line.fill.background()
    circ2 = slide.shapes.add_shape(9, Inches(10.2), Inches(-0.8), Inches(3.5), Inches(3.5))
    circ2.fill.solid(); circ2.fill.fore_color.rgb = _rgb(p["bg_dark"]); circ2.line.fill.background()
    dot = slide.shapes.add_shape(9, Inches(0.6), Inches(6.7), Inches(0.25), Inches(0.25))
    dot.fill.solid(); dot.fill.fore_color.rgb = _rgb(p["accent"]); dot.line.fill.background()
    _add_text_box(slide, "PRESENTATION", Inches(0.6), Inches(1.4), Inches(6), Inches(0.4),
                  "Trebuchet MS", 9, False, p["accent"])
    tx = slide.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(9.5), Inches(2.8))
    tx.text_frame.word_wrap = True
    pr = tx.text_frame.paragraphs[0]; pr.alignment = PP_ALIGN.LEFT
    run = pr.add_run()
    run.text = title; run.font.name = "Georgia"; run.font.size = Pt(48)
    run.font.bold = True; run.font.color.rgb = _rgb(p["text_dark"])
    _add_text_box(slide, "SlideCraft AI  —  AI-Powered Presentation",
                  Inches(0.6), Inches(5.0), Inches(9.5), Inches(0.6),
                  "Calibri", 16, False, p["shape"])
    _add_rect(slide, Inches(0.6), Inches(6.6), Inches(4), Inches(0.04), p["accent"])
    _add_text_box(slide, "Generated by SlideCraft AI", Inches(0.6), Inches(6.75),
                  Inches(6), Inches(0.4), "Calibri", 10, False, "555577")


def _build_layout_a(slide, title: str, bullets: list[str]):
    p = STATE.palette
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, p["bg_light"])
    _add_rect(slide, 0, 0, SLIDE_W, Inches(0.07), p["accent"])
    _add_rect(slide, 0, 0, Inches(0.55), SLIDE_H, p["bg_dark"])
    nb = slide.shapes.add_textbox(Inches(0.0), Inches(3.2), Inches(0.55), Inches(0.5))
    nb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    nr = nb.text_frame.paragraphs[0].add_run()
    nr.text = f"{STATE.slide_index:02d}"; nr.font.name = "Georgia"; nr.font.size = Pt(12)
    nr.font.bold = True; nr.font.color.rgb = _rgb(p["accent"])
    _add_rect(slide, Inches(0.55), Inches(0.07), SLIDE_W - Inches(0.55), Inches(1.3), p["shape"])
    tx = slide.shapes.add_textbox(Inches(0.85), Inches(0.2), Inches(10.5), Inches(1.0))
    tx.text_frame.word_wrap = True
    pr = tx.text_frame.paragraphs[0]; pr.alignment = PP_ALIGN.LEFT
    run = pr.add_run()
    run.text = title; run.font.name = "Georgia"; run.font.size = Pt(32)
    run.font.bold = True; run.font.color.rgb = _rgb(p["text_light"])
    for i, bullet in enumerate(bullets[:6]):
        top = Inches(1.6) + i * Inches(0.85)
        dot = slide.shapes.add_shape(9, Inches(0.85), top + Inches(0.18), Inches(0.15), Inches(0.15))
        dot.fill.solid(); dot.fill.fore_color.rgb = _rgb(p["accent"]); dot.line.fill.background()
        txt = slide.shapes.add_textbox(Inches(1.15), top, Inches(10.0), Inches(0.75))
        txt.text_frame.word_wrap = True
        pr2 = txt.text_frame.paragraphs[0]; pr2.alignment = PP_ALIGN.LEFT
        run2 = pr2.add_run()
        run2.text = bullet; run2.font.name = "Calibri"
        run2.font.size = Pt(18); run2.font.bold = (i == 0)
        run2.font.color.rgb = _rgb(p["text_light"])


def _build_layout_b(slide, title: str, bullets: list[str]):
    p = STATE.palette
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, p["bg_light"])
    _add_rect(slide, Inches(8.0), 0, Inches(5.333), SLIDE_H, p["bg_dark"])
    circ = slide.shapes.add_shape(9, Inches(8.8), Inches(1.5), Inches(3.6), Inches(3.6))
    circ.fill.solid(); circ.fill.fore_color.rgb = _rgb(p["accent"]); circ.line.fill.background()
    circ2 = slide.shapes.add_shape(9, Inches(9.3), Inches(2.0), Inches(2.6), Inches(2.6))
    circ2.fill.solid(); circ2.fill.fore_color.rgb = _rgb(p["bg_dark"]); circ2.line.fill.background()
    _add_text_box(slide, f"{STATE.slide_index:02d}", Inches(9.8), Inches(5.8), Inches(2),
                  Inches(0.6), "Georgia", 28, True, p["accent"], PP_ALIGN.CENTER)
    _add_rect(slide, 0, 0, Inches(8.0), Inches(0.07), p["accent"])
    tx = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(7.2), Inches(1.2))
    tx.text_frame.word_wrap = True
    pr = tx.text_frame.paragraphs[0]; pr.alignment = PP_ALIGN.LEFT
    run = pr.add_run()
    run.text = title; run.font.name = "Georgia"; run.font.size = Pt(30)
    run.font.bold = True; run.font.color.rgb = _rgb(p["text_light"])
    _add_rect(slide, Inches(0.5), Inches(1.55), Inches(2.5), Inches(0.05), p["accent"])
    for i, bullet in enumerate(bullets[:5]):
        top = Inches(1.8) + i * Inches(1.0)
        card = slide.shapes.add_shape(1, Inches(0.5), top, Inches(7.1), Inches(0.8))
        card.fill.solid(); card.fill.fore_color.rgb = _rgb(p["shape"]); card.line.fill.background()
        tb = card.text_frame; tb.word_wrap = True
        pr2 = tb.paragraphs[0]; pr2.alignment = PP_ALIGN.LEFT
        run2 = pr2.add_run()
        run2.text = f"  {bullet}"; run2.font.name = "Calibri"
        run2.font.size = Pt(16); run2.font.color.rgb = _rgb(p["text_light"])


def _build_layout_c(slide, title: str, bullets: list[str]):
    p = STATE.palette
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, p["bg_dark"])
    _add_rect(slide, 0, 0, SLIDE_W, Inches(0.12), p["accent"])
    _add_rect(slide, 0, Inches(0.12), SLIDE_W, Inches(0.04), p["shape"])
    tx = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(11), Inches(1.0))
    tx.text_frame.word_wrap = True
    pr = tx.text_frame.paragraphs[0]; pr.alignment = PP_ALIGN.LEFT
    run = pr.add_run()
    run.text = title; run.font.name = "Georgia"; run.font.size = Pt(34)
    run.font.bold = True; run.font.color.rgb = _rgb(p["text_dark"])
    cols = 3 if len(bullets) >= 4 else 2
    col_w = Inches(3.8) if cols == 3 else Inches(5.5)
    col_gap = Inches(0.35) if cols == 3 else Inches(0.5)
    start_x = Inches(0.7)
    for i, bullet in enumerate(bullets[:6]):
        col = i % cols; row = i // cols
        left = start_x + col * (col_w + col_gap)
        top = Inches(1.7) + row * Inches(2.1)
        card = slide.shapes.add_shape(1, left, top, col_w, Inches(1.7))
        card.fill.solid()
        card.fill.fore_color.rgb = _rgb("1A1A2E" if p["bg_dark"] != "1A1A2E" else "0F0F1E")
        card.line.color.rgb = _rgb(p["accent"]); card.line.width = Pt(1)
        nb2 = slide.shapes.add_shape(9, left + Inches(0.15), top + Inches(0.15), Inches(0.4), Inches(0.4))
        nb2.fill.solid(); nb2.fill.fore_color.rgb = _rgb(p["accent"]); nb2.line.fill.background()
        nbp = nb2.text_frame.paragraphs[0]; nbp.alignment = PP_ALIGN.CENTER
        nbr = nbp.add_run()
        nbr.text = str(i + 1); nbr.font.name = "Trebuchet MS"; nbr.font.size = Pt(10)
        nbr.font.bold = True; nbr.font.color.rgb = _rgb("FFFFFF")
        tb2 = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.65),
                                       col_w - Inches(0.35), Inches(0.95))
        tb2.text_frame.word_wrap = True
        pr2 = tb2.text_frame.paragraphs[0]; pr2.alignment = PP_ALIGN.LEFT
        run2 = pr2.add_run()
        run2.text = bullet; run2.font.name = "Calibri"
        run2.font.size = Pt(14); run2.font.color.rgb = _rgb(p["text_dark"])


def _build_conclusion_slide(slide, title: str, bullets: list[str]):
    p = STATE.palette
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, p["bg_dark"])
    for sx, sy, sr in [(10.5, -0.5, 4.5), (11.8, 1.0, 3.0), (-0.5, 5.5, 3.0)]:
        c = slide.shapes.add_shape(9, Inches(sx), Inches(sy), Inches(sr), Inches(sr))
        c.fill.solid(); c.fill.fore_color.rgb = _rgb(p["accent"]); c.line.fill.background()
        c2 = slide.shapes.add_shape(9, Inches(sx + 0.4), Inches(sy + 0.4),
                                    Inches(sr - 0.8), Inches(sr - 0.8))
        c2.fill.solid(); c2.fill.fore_color.rgb = _rgb(p["bg_dark"]); c2.line.fill.background()
    _add_rect(slide, Inches(5.5), Inches(0.5), Inches(2.3), Inches(0.07), p["accent"])
    tx = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.2))
    tx.text_frame.word_wrap = True
    pr = tx.text_frame.paragraphs[0]; pr.alignment = PP_ALIGN.CENTER
    run = pr.add_run()
    run.text = title; run.font.name = "Georgia"; run.font.size = Pt(38)
    run.font.bold = True; run.font.color.rgb = _rgb(p["text_dark"])
    for i, bullet in enumerate(bullets[:5]):
        top = Inches(2.2) + i * Inches(0.9)
        txt = slide.shapes.add_textbox(Inches(2.0), top, Inches(9.3), Inches(0.75))
        txt.text_frame.word_wrap = True
        pr2 = txt.text_frame.paragraphs[0]; pr2.alignment = PP_ALIGN.CENTER
        run2 = pr2.add_run()
        run2.text = bullet; run2.font.name = "Calibri"
        run2.font.size = Pt(18); run2.font.color.rgb = _rgb(p["shape"])
    _add_rect(slide, Inches(5.5), Inches(7.1), Inches(2.3), Inches(0.07), p["accent"])
    _add_text_box(slide, "Generated by SlideCraft AI", Inches(0.5), Inches(7.15),
                  Inches(12.3), Inches(0.35), "Calibri", 9, False, "555577", PP_ALIGN.CENTER)


# ── MCP tool definitions ───────────────────────────────────────────────────────

@mcp.tool()
def create_presentation() -> str:
    """Initialize a new themed presentation."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    STATE.presentation = prs
    STATE.slides_added = 0
    STATE.slide_index = 0
    STATE.palette = random.choice(PALETTES)
    return "Presentation created (SlideCraft AI themed)"


@mcp.tool()
def add_slide(title: str, bullets: list[str] | None = None) -> str:
    """Add a beautifully themed content slide."""
    prs = _require_presentation()
    normalized = _normalize_bullets(bullets)
    if not normalized:
        normalized = ["Key concept overview", "Important details and examples", "Real-world applications"]

    STATE.slide_index += 1
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    conclusion_titles = {"summary", "conclusion", "q&a", "q&a / references", "thank you", "references"}
    if STATE.slide_index == 1:
        _build_title_slide(slide, title)
    elif title.lower() in conclusion_titles:
        _build_conclusion_slide(slide, title, normalized)
    else:
        layout = STATE.slide_index % 3
        if layout == 0:
            _build_layout_a(slide, title, normalized)
        elif layout == 1:
            _build_layout_b(slide, title, normalized)
        else:
            _build_layout_c(slide, title, normalized)

    STATE.slides_added += 1
    return f"Added slide: {title}"


@mcp.tool()
def add_image_slide(title: str, image_prompt: str = "", caption: str = "") -> str:
    """Add a slide with an AI-generated image or rich decorative placeholder."""
    prs = _require_presentation()
    STATE.slide_index += 1
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    p = STATE.palette

    prompt = image_prompt.strip() or f"High quality educational illustration of {title}"
    image_bytes = _call_hf_image_model(prompt)

    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, p["bg_dark"])
    if image_bytes:
        stream = _fit_image(image_bytes)
        slide.shapes.add_picture(stream, Inches(0.5), Inches(1.1),
                                 width=Inches(12.3), height=Inches(5.5))
        image_status = "AI-generated"
    else:
        _add_rect(slide, Inches(0.5), Inches(1.1), Inches(12.3), Inches(5.5), "1A1A2E")
        for i in range(5):
            sz = Inches(1.5 + i * 0.4)
            circ = slide.shapes.add_shape(9, Inches(4.5 + i * 0.8), Inches(2.0 + i * 0.3), sz, sz)
            circ.fill.solid(); circ.fill.fore_color.rgb = _rgb(p["accent"]); circ.line.fill.background()
        _add_text_box(slide, "[ AI Image — Generation Unavailable ]",
                      Inches(3.5), Inches(3.5), Inches(6), Inches(0.6),
                      "Trebuchet MS", 14, False, p["shape"], PP_ALIGN.CENTER)
        image_status = "placeholder"

    _add_rect(slide, 0, 0, SLIDE_W, Inches(0.95), p["bg_dark"])
    _add_rect(slide, 0, 0, Inches(0.12), SLIDE_H, p["accent"])
    _add_text_box(slide, title, Inches(0.4), Inches(0.12), Inches(12.5), Inches(0.75),
                  "Georgia", 28, True, p["text_dark"])
    _add_rect(slide, 0, Inches(6.65), SLIDE_W, Inches(0.85), p["bg_dark"])
    if caption:
        _add_text_box(slide, caption, Inches(0.5), Inches(6.7), Inches(12.0), Inches(0.65),
                      "Calibri", 13, False, p["shape"])

    STATE.slides_added += 1
    return f"Added image slide: {title} ({image_status})"


@mcp.tool()
def save_presentation(filename: str = "output.pptx") -> str:
    """Save the active presentation to the outputs folder."""
    prs = _require_presentation()
    path = _output_path(filename)
    prs.save(path)
    return f"Saved presentation: {path.name}"


@mcp.tool()
def plan_slides(topic: str, num_slides: int = 5) -> str:
    """Generate a JSON slide plan as a list of {title, bullets}."""
    num_slides = max(2, min(int(num_slides), 15))
    prompt = (
        "Create a slide plan as strict JSON. "
        "Return only a JSON array of objects with keys: title, bullets. "
        f"Topic: {topic}. Number of slides: {num_slides}. "
        "Rules: first slide is title slide with empty bullets; each next slide has 3-5 concise bullets."
    )
    generated = _call_hf_text_model(prompt)
    parsed = _extract_json(generated) if generated else None
    if isinstance(parsed, list) and parsed:
        cleaned = []
        for item in parsed:
            if isinstance(item, dict) and "title" in item:
                cleaned.append({"title": str(item.get("title", "Untitled")),
                                 "bullets": _normalize_bullets(item.get("bullets", []))})
            elif isinstance(item, str):
                cleaned.append({"title": item, "bullets": []})
        if cleaned:
            return json.dumps(cleaned[:num_slides], ensure_ascii=True)
    return json.dumps(build_fallback_plan(topic, num_slides), ensure_ascii=True)


@mcp.tool()
def generate_content(title: str) -> str:
    """Generate JSON bullet points for a single slide title."""
    prompt = (
        "Generate 4 concise presentation bullets as strict JSON array of strings. "
        f"Title: {title}. Rules: under 15 words each, no numbering, no markdown."
    )
    generated = _call_hf_text_model(prompt)
    parsed = _extract_json(generated) if generated else None
    if isinstance(parsed, list) and parsed:
        bullets = [str(x).strip() for x in parsed if str(x).strip()]
        if bullets:
            return json.dumps(bullets[:5], ensure_ascii=True)
    fallback = [f"What {title} is", f"Why {title} matters",
                f"Key components of {title}", f"Practical examples of {title}"]
    return json.dumps(fallback, ensure_ascii=True)


if __name__ == "__main__":
    mcp.run(transport="stdio")